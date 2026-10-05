import json, os, zipfile, io, re
from datetime import datetime, date, timezone
from openpyxl import load_workbook

OUT = os.environ.get("OUTPUT_JSON", "data.json")

EXCEL_FILE = os.environ.get("EXCEL_FILE", "HSE Dashboard Rev 0(3).xlsx")

def download_xlsx(_url=None):
    """Read the dashboard workbook directly from the GitHub repository."""
    if not os.path.exists(EXCEL_FILE):
        raise RuntimeError(f"Excel workbook not found: {EXCEL_FILE}")
    with open(EXCEL_FILE, "rb") as f:
        blob = f.read()
    if len(blob) < 1000 or blob[:2] != b"PK":
        raise RuntimeError(f"{EXCEL_FILE} is not a valid XLSX file ({len(blob)} bytes).")
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as z:
            names = set(z.namelist())
            if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
                raise RuntimeError(f"{EXCEL_FILE} is not a valid XLSX workbook.")
    except zipfile.BadZipFile as e:
        raise RuntimeError(f"{EXCEL_FILE} is not a valid XLSX ZIP file.") from e
    print(f"Loaded workbook from GitHub repository: {EXCEL_FILE} ({len(blob):,} bytes)")
    return blob

def norm(v):
    return re.sub(r"\s+", " ", str(v if v is not None else "").strip()).upper()

def excel_date(v):
    if v is None or v == "":
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, (int,float)):
        # Excel serial date, 1899-12-30 system.
        from datetime import timedelta
        return (datetime(1899,12,30) + timedelta(days=float(v))).strftime("%Y-%m-%d")
    s = str(v).strip()
    m = re.match(r"^(\d{4})[/-](\d{1,2})[/-](\d{1,2})", s)
    if m:
        return f"{m.group(1)}-{int(m.group(2)):02d}-{int(m.group(3)):02d}"
    m = re.match(r"^(\d{1,2})[/-](\d{1,2})[/-](\d{4})", s)
    if m:
        a,b,y = int(m.group(1)),int(m.group(2)),m.group(3)
        if a > 12:
            day,month = a,b
        else:
            month,day = a,b
        if 1 <= month <= 12 and 1 <= day <= 31:
            return f"{y}-{month:02d}-{day:02d}"
    return s[:10]

def week_label(d):
    """Match the workbook WEEK formula: weeks start 11/04/2026 (W01)."""
    if not d:
        return ""
    try:
        base=date(2026,4,11)
        cur=date.fromisoformat(str(d)[:10])
        if cur < base:
            return "0"
        w=((cur-base).days//7)+1
        start=base.fromordinal(base.toordinal()+(w-1)*7)
        end=start.fromordinal(start.toordinal()+6)
        return f"W{w:02d} ({start.day}/{start.month} - {end.day}/{end.month})"
    except Exception:
        return ""

def hazard_label(raw, bsc_type):
    """Match the workbook Hazard Type formula."""
    raw=str(raw or "").strip()
    bsc_type=str(bsc_type or "").strip()
    if raw.upper()=="NA" or not raw:
        return bsc_type
    if raw.lower()=="others":
        return f"Others - {bsc_type}" if bsc_type else "Others"
    return raw

def num(v):
    if isinstance(v,(int,float)) and not isinstance(v,bool):
        return float(v)
    try:
        return float(str(v if v is not None else "").replace(",",""))
    except:
        return 0

def rows_from_sheet(ws):
    rows = list(ws.iter_rows(values_only=True))
    hr = None
    for i,row in enumerate(rows):
        vals = [norm(x) for x in row]
        if "DATE" in vals or "DATE REPORTED" in vals:
            hr = i
            break
    if hr is None:
        return []
    headers = [norm(x) if norm(x) else f"__EMPTY_{i}" for i,x in enumerate(rows[hr])]
    out=[]
    for row in rows[hr+1:]:
        obj={}
        has=False
        for i,h in enumerate(headers):
            v=row[i] if i < len(row) else ""
            obj[h]=v
            if v not in ("",None):
                has=True
        if has:
            out.append(obj)
    return out

def val(row,names):
    names = names if isinstance(names,list) else [names]
    for n in names:
        if norm(n) in row:
            return row[norm(n)]
    return ""

def clean_sheet_name(s):
    return str(s)

wb = load_workbook(EXCEL_FILE, data_only=True, read_only=True)

def get_sheet(*names):
    """Find a worksheet by name, ignoring case, extra spaces and minor naming differences."""
    targets = {norm(n).replace("_", " ") for n in names}
    for ws in wb.worksheets:
        actual = norm(ws.title).replace("_", " ")
        if actual in targets:
            return ws
    return None

excluded={"LAGGING","LEADING","BE SAFE CARD LOG","TRAINING"}
manpower=[]
for ws in wb.worksheets:
    if norm(ws.title) in excluded:
        continue
    raw=rows_from_sheet(ws)
    if not raw:
        continue
    first=raw[0]
    if "DATE" not in first or not ("MANHOUR" in first or "MANDAY" in first):
        continue
    rows=[]
    for r in raw:
        d=excel_date(val(r,"DATE"))
        if not d:
            continue
        rows.append({
            "date":d,
            "week":str(val(r,"WEEK") or "") or week_label(d),
            "company":str(val(r,"COMPANY") or ""),
            "category":str(val(r,"CATEGORY") or ""),
            "manday":num(val(r,"MANDAY")),
            "manhour":num(val(r,"MANHOUR"))
        })
    if rows:
        manpower.append({"sheet":clean_sheet_name(ws.title),"rows":rows})

bsc_ws=get_sheet("Be Safe Card Log", "BE SAFE CARD LOG", "BESAFE CARD LOG")
bsc_raw=rows_from_sheet(bsc_ws) if bsc_ws else []
bsc=[]
for i,r in enumerate(bsc_raw,1):
    bsc.append({
        "no":num(val(r,"NO")) or i,
        "bscNo":str(val(r,["BSC NO","BSC NO."]) or ""),
        "type":str(val(r,["BSC TYPE","TYPE"]) or ""),
        "hazardRaw":str(val(r,"HAZARD TYPE RAW") or ""),
        "hazard":hazard_label(val(r,"HAZARD TYPE RAW"), val(r,["BSC TYPE","TYPE"])),
        "dateReported":excel_date(val(r,"DATE REPORTED")),
        "week":str(val(r,"WEEK") or "") or week_label(excel_date(val(r,"DATE REPORTED"))),
        "dateObserved":excel_date(val(r,"DATE OBSERVED")),
        "observer":str(val(r,["OBSERVED BY","OBSERVER"]) or ""),
        "company":str(val(r,"COMPANY") or ""),
        "position":str(val(r,"POSITION") or ""),
        "description":str(val(r,"WHAT AND WHERE HAZARD/ACT OBSERVED AND DISCUSSED WITH WORKER/SUPERVISOR/MANAGER?") or ""),
        "status":str(val(r,"STATUS") or "")
    })
bsc=[r for r in bsc if r["dateReported"] or r["dateObserved"]]

tr_topics=["SIC","PTW NOVADE","EMERGENCY RESPONSE","FIRE WATCH","CHEMICAL HANDLING","WASTE MANAGEMENT","ELECTRICAL SAFETY","PPE AWARENESS","HAND TOOLS & POWER TOOLS SAFETY"]
tr_ws=get_sheet("Training")
tr_raw=rows_from_sheet(tr_ws) if tr_ws else []
training=[]
for r in tr_raw:
    d=excel_date(val(r,"DATE"))
    if d:
        o={"date":d,"week":str(val(r,"WEEK") or "") or week_label(d)}
        for t in tr_topics: o[t]=num(val(r,t))
        training.append(o)

lead_topics=["HSE INDUCTION","TOOLBOX MEETING","HSE COORDINATION MEETING","HSE COMMITTEE MEETING","EMERGENCY DRILL","WORKPLACE INSPECTION","HSE WALKABOUT","WEEKLY HOUSEKEEPING","BESAFE CARD","HSE CAMPAIGN","REWARDS & RECOGNITION","CONSEQUENCE MANAGEMENT"]
lead_ws=get_sheet("Leading")
lead_raw=rows_from_sheet(lead_ws) if lead_ws else []
leading=[]
for r in lead_raw:
    d=excel_date(val(r,"DATE"))
    if d:
        o={"date":d,"week":str(val(r,"WEEK") or "") or week_label(d)}
        for t in lead_topics: o[t]=num(val(r,t))
        leading.append(o)

lag_topics=["FATALITY","LTI","RWC","MTC","FAC","DANGEROUS OCCURANCE","PROPERTY DAMAGE","NEAR MISS","FIRE INCIDENT","REGULATORY NON COMPLIANCE","LoPC"]
lag_ws=get_sheet("Lagging")
lag_raw=rows_from_sheet(lag_ws) if lag_ws else []

users_ws=get_sheet("Users")
users=[]
if users_ws:
    user_rows=list(users_ws.iter_rows(values_only=True))
    if user_rows:
        user_headers=[norm(x) if norm(x) else f"__EMPTY_{i}" for i,x in enumerate(user_rows[0])]
        user_index={h:i for i,h in enumerate(user_headers)}
        for row in user_rows[1:]:
            def user_val(name):
                idx=user_index.get(norm(name))
                return row[idx] if idx is not None and idx < len(row) else ""
            username=str(user_val("USERNAME") or "").strip()
            if username:
                users.append({
                    "username":username,
                    "password":str(user_val("PASSWORD") or ""),
                    "name":str(user_val("NAME") or ""),
                    "role":str(user_val("ROLE") or ""),
                    "status":str(user_val("STATUS") or "")
                })

lagging=[]
for r in lag_raw:
    d=excel_date(val(r,"DATE"))
    if d:
        o={"date":d,"week":str(val(r,"WEEK") or "") or week_label(d)}
        for t in lag_topics: o[t]=num(val(r,t))
        lagging.append(o)

payload={
    "generated_at":datetime.now(timezone.utc).isoformat(),
    "source":"GitHub repository",
    "workbook":os.path.basename(EXCEL_FILE),
    "manpower":manpower,
    "bsc":bsc,
    "training":training,
    "trainingTopics":tr_topics,
    "leading":leading,
    "leadingTopics":lead_topics,
    "lagging":lagging,
    "laggingTopics":lag_topics,
    "users":users
}
with open(OUT,"w",encoding="utf-8") as f:
    json.dump(payload,f,ensure_ascii=False,separators=(",",":"))
print("Generated",OUT)
print("Manpower rows:",sum(len(x["rows"]) for x in manpower))
print("BSC:",len(bsc),"Training:",len(training),"Leading:",len(leading),"Lagging:",len(lagging))
